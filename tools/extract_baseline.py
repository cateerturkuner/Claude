"""Extract a post-close proforma baseline from a Walters Hospitality venue
forecast workbook (Detailed P&L + Revenue Analysis tabs) into JSON.

Usage:  python3 tools/extract_baseline.py <workbook.xlsx> <venue-slug> [close YYYY-MM]
"""
import json
import sys
from datetime import datetime

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# ---------------------------------------------------------------- Detailed P&L
# Row map for the Detailed P&L tab. Order matters -- it is the display order.
PL_ROWS = [
    ("funnel", "leads", 11, "Leads", "count"),
    ("funnel", "lead_to_tour", 13, "Lead → Tour Conversion", "pct"),
    ("funnel", "tours", 15, "Tours", "count"),
    ("funnel", "tour_to_contract", 17, "Tour → Contract Conversion", "pct"),
    ("funnel", "contracts", 19, "Contracts", "count"),
    ("funnel", "events_oc", 21, "Events – Original Contracts", "count"),
    ("funnel", "events_new", 22, "Events – New Contracts", "count"),
    ("funnel", "events", 23, "Events", "count"),

    ("revenue", "room_rental", 27, "Room Rental Revenue", "usd"),
    ("revenue", "food", 28, "Food Revenue", "usd"),
    ("revenue", "beverage", 29, "Beverage Revenue", "usd"),
    ("revenue", "lodging", 30, "Lodging Revenues", "usd"),
    ("revenue", "ancillary", 31, "Ancillary Revenue", "usd"),
    ("revenue", "service_charge", 32, "Service Charge", "usd"),
    ("revenue", "cancelled_events", 33, "Cancelled Events", "usd"),
    ("revenue", "outside_events", 34, "Outside Events", "usd"),
    ("revenue", "discounts", 35, "Discounts", "usd"),
    ("revenue", "other_revenue", 36, "Other Revenue", "usd"),
    ("total", "total_revenue", 38, "Total Revenue", "usd"),

    ("cogs", "food_cogs", 41, "Food COGS", "usd"),
    ("cogs", "ancillary_cogs", 42, "Ancillary COGS", "usd"),
    ("cogs", "alcohol_cogs", 43, "Alcohol COGS", "usd"),
    ("cogs", "other_cogs", 44, "Other COGS", "usd"),
    ("cogs", "room_cogs", 45, "Room COGS", "usd"),
    ("total", "total_cogs", 47, "Total COGS", "usd"),
    ("total", "gross_margin", 49, "Gross Margin", "usd"),

    ("payroll", "sales_payroll", 52, "Sales Team Payroll", "usd"),
    ("payroll", "planning_payroll", 53, "Planning Team Payroll", "usd"),
    ("payroll", "operations_payroll", 54, "Operations Team Payroll", "usd"),
    ("payroll", "ancillary_payroll", 55, "Ancillary Payroll", "usd"),
    ("payroll", "fnb_payroll", 56, "F&B Team Payroll", "usd"),
    ("payroll", "owner_salaries", 57, "Owner Salaries", "usd"),
    ("payroll", "payroll_admin", 58, "Payroll Admin & Benefit Expenses", "usd"),
    ("payroll", "other_payroll", 59, "Other Payroll", "usd"),
    ("total", "total_payroll", 61, "Total Payroll", "usd"),

    ("opex", "marketing", 64, "Marketing Expense", "usd"),
    ("opex", "maintenance", 65, "Maintenance Expense", "usd"),
    ("opex", "utilities", 66, "Utilities Expense", "usd"),
    ("opex", "office", 67, "Office Expense", "usd"),
    ("opex", "computer", 68, "Computer Expense", "usd"),
    ("opex", "automobile", 69, "Automobile Expenses", "usd"),
    ("opex", "training", 70, "Training & Development Expense", "usd"),
    ("opex", "tax", 71, "Tax Expense", "usd"),
    ("opex", "travel", 72, "Travel Expense", "usd"),
    ("opex", "professional", 73, "Professional Expenses", "usd"),
    ("opex", "rent_expense", 74, "Rent Expense", "usd"),
    ("opex", "insurance", 75, "Insurance", "usd"),
    ("opex", "finance", 76, "Finance Expenses", "usd"),
    ("opex", "personal", 77, "Personal Expenses", "usd"),
    ("opex", "other_opex", 78, "Other Operating Expenses", "usd"),
    ("total", "total_opex", 80, "Total Operating Expenses", "usd"),

    ("total", "ebitdar", 82, "EBITDAR", "usd"),
    ("total", "rent", 84, "Rent", "usd"),
    ("total", "ebitda", 86, "EBITDA", "usd"),
]

# Rows 24/25 hold the monthly share of the post-close year's events. Row 25 is
# the smoothed seasonality actually used to spread the annual event count.
EVENT_SEASONALITY_ROW = 25

# ----------------------------------------------------------- Revenue Analysis
# Ancillary / revenue categories, with the Revenue Analysis row that holds the
# attachment count (events) and the row that holds the revenue for each of the
# Original-Contract (OC) and New-Contract (NC) blocks.
RA_CATEGORIES = [
    ("room_rental", "Room Rental", 29, 43, 60, 75),
    ("food", "Food", 30, 44, 61, 76),
    ("beverage", "Beverage", 31, 45, 62, 77),
    ("lodging", "Lodging", 32, 46, 63, 78),
    ("dj", "DJ", 33, 47, 64, 79),
    ("floral", "Floral", 34, 48, 65, 80),
    ("bakery", "Bakery", 35, 49, 66, 81),
    ("stationery", "Stationery", 36, 50, 67, 82),
    ("photography", "Photography", 37, 51, 68, 83),
    ("other_ancillary", "Other Ancillary", 38, 52, 69, 84),
]
RA_SERVICE_CHARGE_NC_REV_ROW = 85

RA_YEAR_COLS = {1: "G", 2: "H", 3: "I"}      # absolute Y1/Y2/Y3 values
RA_DRIVER_COLS = {1: "K", 2: "L", 3: "M"}    # attachment % or avg $/event
RA_COMMENT_COL = "T"
RA_BENCH_COLS = {"venue": "P", "region": "Q", "wh_avg": "R"}


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    return None


def cell(ws, row, col_letter):
    return num(ws[f"{col_letter}{row}"].value)


def extract(path, slug, close_ym=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    pl = wb["Detailed P&L"]
    ra = wb["Revenue Analysis"]

    venue_name = pl["C5"].value or slug

    # --- locate the first post-close forecast column -----------------------
    # Row 3 carries the post-close year index (1, 2, 3...) for forecast months.
    first_col = None
    for c in range(1, pl.max_column + 1):
        if isinstance(pl.cell(3, c).value, (int, float)) and isinstance(
            pl.cell(9, c).value, datetime
        ):
            first_col = c
            break
    if first_col is None:
        raise SystemExit("Could not find the forecast block (row 3 year tags).")

    if close_ym:
        want = datetime.strptime(close_ym, "%Y-%m")
        for c in range(first_col, pl.max_column + 1):
            d = pl.cell(9, c).value
            if isinstance(d, datetime) and (d.year, d.month) == (want.year, want.month):
                first_col = c
                break

    months = []
    for c in range(first_col, pl.max_column + 1):
        d = pl.cell(9, c).value
        py = pl.cell(3, c).value
        if not isinstance(d, datetime) or not isinstance(py, (int, float)):
            break
        months.append({
            "col": get_column_letter(c),
            "ym": f"{d.year:04d}-{d.month:02d}",
            "date": d.strftime("%Y-%m-%d"),
            "post_close_year": int(py),
            "event_share": num(pl.cell(EVENT_SEASONALITY_ROW, c).value),
            "lines": {},
        })

    lines_meta = []
    for group, key, row, label, fmt in PL_ROWS:
        lines_meta.append({"group": group, "key": key, "label": label, "format": fmt,
                           "source_row": row})
        for m in months:
            m["lines"][key] = num(pl[f"{m['col']}{row}"].value)

    for m in months:
        m.pop("col")

    # --- Revenue Analysis: annual attachment + $/event drivers -------------
    drivers = {"years": {}, "categories": [], "benchmarks": {}}

    for year, gcol in RA_YEAR_COLS.items():
        drivers["years"][str(year)] = {
            "events_oc": cell(ra, 11, gcol),
            "events_new": cell(ra, 12, gcol),
            "events_total": cell(ra, 13, gcol),
            "revenue_oc": cell(ra, 16, gcol),
            "revenue_new": cell(ra, 18, gcol),
            "revenue_total": cell(ra, 19, gcol),
            "rev_per_event_oc": cell(ra, 22, gcol),
            "rev_per_event_new": cell(ra, 24, gcol),
            "rev_per_event_total": cell(ra, 25, gcol),
        }

    for key, label, oc_att_r, nc_att_r, oc_rev_r, nc_rev_r in RA_CATEGORIES:
        entry = {
            "key": key,
            "label": label,
            "comment": (ra[f"{RA_COMMENT_COL}{nc_att_r}"].value
                        or ra[f"{RA_COMMENT_COL}{nc_rev_r}"].value or None),
            "benchmark": {
                name: cell(ra, nc_rev_r, col)
                for name, col in RA_BENCH_COLS.items()
            },
            "oc": {"attach_rate": {}, "attach_events": {}, "revenue": {},
                   "rev_per_event": {}},
            "new": {"attach_rate": {}, "attach_events": {}, "revenue": {},
                    "rev_per_event": {}},
        }
        for year in (1, 2, 3):
            y, gcol, dcol = str(year), RA_YEAR_COLS[year], RA_DRIVER_COLS[year]
            entry["oc"]["attach_rate"][y] = cell(ra, oc_att_r, dcol)
            entry["oc"]["attach_events"][y] = cell(ra, oc_att_r, gcol)
            entry["oc"]["revenue"][y] = cell(ra, oc_rev_r, gcol)
            entry["oc"]["rev_per_event"][y] = cell(ra, oc_rev_r, dcol)
            entry["new"]["attach_rate"][y] = cell(ra, nc_att_r, dcol)
            entry["new"]["attach_events"][y] = cell(ra, nc_att_r, gcol)
            entry["new"]["revenue"][y] = cell(ra, nc_rev_r, gcol)
            entry["new"]["rev_per_event"][y] = cell(ra, nc_rev_r, dcol)
        drivers["categories"].append(entry)

    drivers["service_charge"] = {
        "label": "Service Charge",
        "comment": ra[f"{RA_COMMENT_COL}{RA_SERVICE_CHARGE_NC_REV_ROW}"].value,
        "new": {
            "rate": {str(y): cell(ra, RA_SERVICE_CHARGE_NC_REV_ROW, RA_DRIVER_COLS[y])
                     for y in (1, 2, 3)},
            "revenue": {str(y): cell(ra, RA_SERVICE_CHARGE_NC_REV_ROW, RA_YEAR_COLS[y])
                        for y in (1, 2, 3)},
        },
    }

    baseline = {
        "slug": slug,
        "venue": venue_name,
        "units": "thousands",
        "source_file": path.split("/")[-1],
        "extracted_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "close_month": months[0]["ym"] if months else None,
        "lines": lines_meta,
        "months": months,
        "drivers": drivers,
    }
    return baseline


if __name__ == "__main__":
    src, slug = sys.argv[1], sys.argv[2]
    close = sys.argv[3] if len(sys.argv) > 3 else None
    data = extract(src, slug, close)
    out = f"postclose/data/baseline_{slug}.json"
    with open(out, "w") as fh:
        json.dump(data, fh, indent=1)
    yrs = {}
    for m in data["months"]:
        yrs.setdefault(m["post_close_year"], 0)
        yrs[m["post_close_year"]] += 1
    print(f"{out}: {len(data['months'])} months from {data['close_month']}, "
          f"months per post-close year {yrs}")
