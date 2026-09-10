"""Import a venue's monthly actuals from a post-close GL extract.

    python3 tools/import_gl_actuals.py <source.xlsx> <venue-slug> [--commit]

The workbook is the one finance produces per venue: a `Source` sheet holding GL
lines, and a `Fields` sheet mapping (Department name, Acct Name, IS Line) onto
the proforma's own line names. That mapping is finance's own work, so it is read
rather than reimplemented -- but it is also verified: any GL row that sits on a
real income-statement line and has no mapping is reported rather than silently
dropped, because a missing mapping would otherwise read as an underspend.

Runs as a dry run by default and prints what it would write. Pass --commit to
store it.
"""
import argparse
import sys
from collections import defaultdict

import openpyxl

sys.path.insert(0, __file__.rsplit("/tools/", 1)[0])
from postclose import analysis, lines as L, store  # noqa: E402

# GL columns, by header name, so a reordered export still works.
WANT = {"location name": "loc", "department name": "dept", "amount": "amt",
        "acct name": "acct", "month": "month", "is line": "isline",
        "entry date": "date", "region": "region"}

# A blank IS Line marks a balance-sheet row. Those are not P&L and are expected
# to be unmapped, so they are excluded rather than reported as a gap.
NON_PL_IS = {"none", "0", "00:00:00", ""}

# Lines the proforma carries but which are not costs of the venue's operations.
# Depreciation is excluded on purpose: EBITDAR is before D&A, and including it
# would understate the line we compare on.
EXCLUDED_ACCOUNTS = {"depreciation expense"}

# EBITDAR is before rent. Finance codes rent into operating expenses, so left
# alone it would sit inside Total Opex and pull EBITDAR down against a proforma
# that carries no rent at all -- for Firefly that is a ~$40k/month error. Routed
# to its own line it stays out of EBITDAR and is still reported, which is the
# useful outcome: rent we are paying that was never underwritten.
RETARGET = {"rent expense": "rent"}


def norm(v):
    return "" if v is None else " ".join(str(v).split()).strip().lower()


def read_fields(wb):
    """(department, account, is line) -> proforma line label, from the Fields sheet."""
    ws = wb["Fields"]
    out = {}
    for r in range(2, ws.max_row + 1):
        dept, acct, isline, target = [ws.cell(r, c).value for c in range(1, 5)]
        if acct is None:
            continue
        out[(norm(dept), norm(acct), norm(isline))] = str(target).strip()
    return out


def label_to_key(slug):
    """Proforma line label -> the baseline's key for it."""
    out = {}
    for ln in analysis.line_meta(slug):
        out[norm(ln["label"])] = ln["key"]
    for key, spec in L.BY_KEY.items():
        out.setdefault(norm(spec["label"]), key)
    return out


def load(path, slug):
    wb = openpyxl.load_workbook(path, data_only=True)
    fields = read_fields(wb)
    keys = label_to_key(slug)
    base_months = {m["ym"]: m for m in store.baseline(slug)["months"]}

    ws = wb["Source"]
    header = {norm(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}
    col = {}
    for name, role in WANT.items():
        if name in header:
            col[role] = header[name]
    missing = {"loc", "dept", "amt", "acct", "isline", "date"} - set(col)
    if missing:
        raise SystemExit(f"Source sheet is missing column(s): {sorted(missing)}")

    monthly = defaultdict(lambda: defaultdict(float))
    unmapped = defaultdict(float)
    excluded = defaultdict(float)
    outside = defaultdict(float)
    locations = set()

    for r in range(2, ws.max_row + 1):
        amt = ws.cell(r, col["amt"]).value
        if amt in (None, ""):
            continue
        try:
            amt = float(amt)
        except (TypeError, ValueError):
            continue
        date = ws.cell(r, col["date"]).value
        if date is None or not hasattr(date, "year"):
            continue
        ym = f"{date.year:04d}-{date.month:02d}"
        locations.add(ws.cell(r, col["loc"]).value)

        dept = norm(ws.cell(r, col["dept"]).value)
        acct = norm(ws.cell(r, col["acct"]).value)
        isline = norm(ws.cell(r, col["isline"]).value)

        if acct in EXCLUDED_ACCOUNTS:
            excluded[acct] += amt
            continue
        if isline in NON_PL_IS:
            continue

        target = fields.get((dept, acct, isline))
        if target is None:
            unmapped[(dept, acct, isline)] += amt
            continue
        if target in ("0", "None", ""):
            continue
        key = RETARGET.get(norm(target)) or keys.get(norm(target))
        if key is None:
            unmapped[(dept, acct, f"{isline} -> unknown PF line '{target}'")] += amt
            continue
        if ym not in base_months:
            outside[ym] += amt
            continue
        # The proforma carries revenue positive and costs negative; the GL is the
        # other way round, so one flip serves every line.
        monthly[ym][key] += -amt / 1000.0

    wb.close()
    return {
        "months": {ym: {k: round(v, 6) for k, v in vals.items()}
                   for ym, vals in sorted(monthly.items())},
        "unmapped": dict(unmapped), "excluded": dict(excluded),
        "outside": dict(outside), "locations": sorted(x for x in locations if x),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("slug")
    ap.add_argument("--commit", action="store_true")
    args = ap.parse_args()

    res = load(args.source, args.slug)
    base = store.baseline(args.slug)
    plan = {m["ym"]: m["lines"] for m in base["months"]}
    rules = analysis.subtotal_rules(base["lines"])
    profit = base["profit_key"]

    print(f"locations in file: {res['locations']}")
    if res["excluded"]:
        print("excluded from EBITDAR by design: "
              + ", ".join(f"{k} {v/1000:+.1f}k" for k, v in res["excluded"].items()))
    if res["outside"]:
        print(f"outside the proforma window, ignored: "
              f"{ {k: round(v/1000,1) for k,v in res['outside'].items()} }")
    if res["unmapped"]:
        print(f"\n!! {len(res['unmapped'])} income-statement row(s) with no mapping "
              f"in Fields — these are NOT imported:")
        for (d, a, i), v in sorted(res["unmapped"].items(), key=lambda x: -abs(x[1])):
            print(f"   {a[:36]:38s} dept={d[:18]:20s} IS={i[:22]:24s} {v/1000:+9.1f}k")
    else:
        print("\nevery income-statement row mapped cleanly")

    print(f"\n{'month':8s} {'revenue':>9} {'COGS':>9} {'payroll':>9} {'opex':>9} "
          f"{profit.upper():>10} {'vs plan':>9} {'rent':>8}")
    for ym, vals in res["months"].items():
        bucket = analysis.fill_subtotals(dict(vals), rules)
        pl = analysis.fill_subtotals({k: v for k, v in plan[ym].items()}, rules)
        var = (bucket.get(profit) or 0) - (pl.get(profit) or 0)
        print(f"{ym:8s} {bucket.get('total_revenue') or 0:9.2f} "
              f"{bucket.get('total_cogs') or 0:9.2f} {bucket.get('total_payroll') or 0:9.2f} "
              f"{bucket.get('total_opex') or 0:9.2f} {bucket.get(profit) or 0:10.2f} "
              f"{var:+9.2f} {vals.get('rent') or 0:8.2f}")

    if not args.commit:
        print("\ndry run — nothing stored. Pass --commit to write.")
        return
    for ym, vals in res["months"].items():
        store.record(args.slug, ym, vals, "upload",
                     note=args.source.split("/")[-1])
    print(f"\nstored {len(res['months'])} month(s) for {args.slug}")


if __name__ == "__main__":
    main()
