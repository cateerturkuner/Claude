"""Read a monthly actuals workbook from finance and map it onto the model.

The finance file is not a fixed template, so nothing here assumes a layout. The
parser looks for a row of month headers and a column of line labels, matches the
labels against the proforma's own line names (plus an alias table), and returns a
preview for a person to confirm before anything is stored.

Two corrections are detected and proposed rather than applied silently:

  scale  a file kept in whole dollars against a model kept in thousands
  sign   cost lines delivered positive against a model that carries them negative
"""
import difflib
import re
from datetime import datetime

from . import analysis, lines as L, store

# Names finance uses that differ from the model's own labels.
ALIASES = {
    "leads": ["leads", "new leads", "inquiries", "inquiry", "enquiries"],
    "tours": ["tours", "tours completed", "site visits", "showings"],
    "contracts": ["contracts", "contracts signed", "bookings", "new bookings",
                  "contracts sold"],
    "events": ["events", "events held", "total events", "weddings and events"],
    "events_oc": ["events original contracts", "oc events", "original contracts",
                  "assumed contracts", "inherited events", "pre close events",
                  "original", "backlog events"],
    "events_new": ["events new contracts", "new events", "new contracts",
                   "new contract events", "new"],
    "room_rental": ["room rental revenue", "venue fee", "venue rental",
                    "rental revenue", "room rental", "venue revenue",
                    "site fee"],
    "food": ["food revenue", "food", "catering revenue", "catering", "f and b food"],
    "beverage": ["beverage revenue", "beverage", "bar revenue", "alcohol revenue",
                 "bar"],
    "lodging": ["lodging revenues", "lodging revenue", "lodging", "accommodation",
                "cabins", "rooms revenue", "overnight"],
    "ancillary": ["ancillary revenue", "ancillary", "add ons", "add on revenue",
                  "vendor revenue", "enhancements"],
    "service_charge": ["service charge", "service fee", "admin fee",
                       "service charges"],
    "cancelled_events": ["cancelled events", "canceled events", "cancellation revenue",
                         "forfeited deposits", "retained deposits"],
    "outside_events": ["outside events", "offsite events", "external events"],
    "discounts": ["discounts", "discount", "concessions", "comps"],
    "other_revenue": ["other revenue", "misc revenue", "miscellaneous revenue",
                      "other income"],
    "total_revenue": ["total revenue", "total sales", "gross revenue", "net revenue",
                      "revenue total"],
    "food_cogs": ["food cogs", "food cost", "cost of food", "catering cogs"],
    "ancillary_cogs": ["ancillary cogs", "ancillary cost", "vendor cost",
                       "cost of ancillary"],
    "alcohol_cogs": ["alcohol cogs", "beverage cogs", "bar cogs", "alcohol cost",
                     "cost of beverage"],
    "other_cogs": ["other cogs", "other cost of sales", "misc cogs"],
    "room_cogs": ["room cogs", "lodging cogs", "lodging cost", "room cost"],
    "total_cogs": ["total cogs", "total cost of goods sold", "cost of goods sold",
                   "total cost of sales"],
    "gross_margin": ["gross margin", "gross profit"],
    "sales_payroll": ["sales team payroll", "sales payroll", "sales wages",
                      "sales salaries"],
    "planning_payroll": ["planning team payroll", "planning payroll",
                         "planning wages", "coordinator payroll", "event planning"],
    "operations_payroll": ["operations team payroll", "operations payroll",
                           "ops payroll", "operations wages", "venue payroll"],
    "ancillary_payroll": ["ancillary payroll", "ancillary wages"],
    "fnb_payroll": ["f and b team payroll", "fnb payroll", "f b payroll",
                    "food and beverage payroll", "kitchen payroll",
                    "banquet payroll"],
    "owner_salaries": ["owner salaries", "owners salary", "owner compensation",
                       "owner draw"],
    "payroll_admin": ["payroll admin and benefit expenses", "payroll taxes",
                      "benefits", "payroll admin", "payroll burden",
                      "taxes and benefits"],
    "other_payroll": ["other payroll", "misc payroll", "contract labor"],
    "total_payroll": ["total payroll", "total labor", "total wages"],
    "marketing": ["marketing expense", "marketing", "advertising", "advertising and marketing"],
    "maintenance": ["maintenance expense", "maintenance", "repairs and maintenance",
                    "r and m", "grounds"],
    "utilities": ["utilities expense", "utilities", "utility"],
    "office": ["office expense", "office", "office supplies", "supplies"],
    "computer": ["computer expense", "computer", "software", "it expense",
                 "technology"],
    "automobile": ["automobile expenses", "automobile", "auto expense", "vehicle",
                   "fuel"],
    "training": ["training and development expense", "training", "development",
                 "training and development"],
    "tax": ["tax expense", "taxes", "property tax", "property taxes"],
    "travel": ["travel expense", "travel", "travel and entertainment"],
    "professional": ["professional expenses", "professional fees", "legal and accounting",
                     "legal", "accounting", "consulting"],
    "rent_expense": ["rent expense", "equipment rent", "equipment rental"],
    "insurance": ["insurance", "insurance expense"],
    "finance": ["finance expenses", "bank fees", "merchant fees", "processing fees",
                "credit card fees", "interest expense"],
    "personal": ["personal expenses", "personal"],
    "other_opex": ["other operating expenses", "other opex", "miscellaneous expense",
                   "other expense", "general and administrative"],
    "total_opex": ["total operating expenses", "total opex", "total expenses",
                   "total operating expense"],
    "ebitdar": ["ebitdar"],
    "rent": ["rent", "base rent", "lease expense", "ground lease"],
    "ebitda": ["ebitda", "net operating income", "noi"],
}

# Lines the model carries negative. Derived from the venue's own line groups
# rather than a fixed list, so a venue with an all-inclusive COGS line or a
# corporate payroll line has them negated too -- a missed cost key silently
# turns an overspend into a favourable variance.
COST_GROUPS = ("cogs", "payroll", "opex")
EXTRA_COST_KEYS = {"rent", "discounts", "total_cogs", "total_payroll",
                   "total_opex"}


def cost_keys(slug):
    keys = {ln["key"] for ln in analysis.line_meta(slug)
            if ln["group"] in COST_GROUPS}
    return keys | EXTRA_COST_KEYS


# Lines that are counts or ratios and must never be rescaled or sign-flipped.
NON_DOLLAR = {"leads", "tours", "contracts", "events", "events_oc", "events_new",
              "lead_to_tour", "tour_to_contract"}

MATCH_THRESHOLD = 0.84
MONTH_NAMES = {m.lower(): i for i, m in enumerate(
    ["", "january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"]) if m}
MONTH_ABBR = {m.lower(): i for i, m in enumerate(analysis.MONTH_ABBR) if m}


def norm(text):
    if text is None:
        return ""
    text = str(text).replace("&", " and ")
    text = re.sub(r"[^a-z0-9 ]+", " ", text.lower())
    return re.sub(r"\s+", " ", text).strip()


def _alias_index(slug):
    idx = {}
    for ln in analysis.line_meta(slug):
        idx.setdefault(norm(ln["label"]), ln["key"])
    for key, names in ALIASES.items():
        for name in names:
            idx.setdefault(norm(name), key)
    # Per-segment contract rows, however the sheet spells the separator.
    for seg in store.baseline(slug).get("segments", []):
        for pattern in ("contracts {}", "{} contracts", "contracts - {}",
                        "contracts \u2013 {}"):
            for name in (seg["label"], seg["key"]):
                idx.setdefault(norm(pattern.format(name)),
                               f"contracts__{seg['key']}")
    return idx


def match_label(raw, idx):
    """Map one spreadsheet label to a model line key, with a confidence."""
    n = norm(raw)
    if not n or len(n) < 2:
        return None, 0.0
    if n in idx:
        return idx[n], 1.0
    best = difflib.get_close_matches(n, list(idx), n=1, cutoff=MATCH_THRESHOLD)
    if best:
        return idx[best[0]], round(difflib.SequenceMatcher(None, n, best[0]).ratio(), 3)
    return None, 0.0


def parse_period(value):
    """Turn a header cell into YYYY-MM, or None if it is not a month."""
    if isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, (int, float)):
        return None
    if not isinstance(value, str):
        return None
    text = value.strip()
    m = re.match(r"^(\d{4})[-/](\d{1,2})$", text)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2})[-/](\d{4})$", text)
    if m:
        return f"{int(m.group(2)):04d}-{int(m.group(1)):02d}"
    m = re.match(r"^([A-Za-z]{3,9})[\s\-']+(\d{2,4})$", text)
    if m:
        name = m.group(1).lower()
        month = MONTH_NAMES.get(name) or MONTH_ABBR.get(name[:3])
        if month:
            year = int(m.group(2))
            year += 2000 if year < 100 else 0
            return f"{year:04d}-{month:02d}"
    return None


def _find_header(ws, max_scan=40):
    """Best (row, {col: ym}) pair of month headers on a sheet."""
    best = (None, {})
    for r in range(1, min(ws.max_row, max_scan) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            ym = parse_period(ws.cell(r, c).value)
            if ym:
                found[c] = ym
        if len(found) > len(best[1]):
            best = (r, found)
    return best


def _find_label_col(ws, header_row, idx, max_col=8):
    """Column of line labels: the one matching the most known model lines."""
    best, best_hits = 1, -1
    for c in range(1, min(ws.max_column, max_col) + 1):
        hits = 0
        for r in range(header_row + 1, ws.max_row + 1):
            key, _ = match_label(ws.cell(r, c).value, idx)
            if key:
                hits += 1
        if hits > best_hits:
            best, best_hits = c, hits
    return best, best_hits


def preview(path, slug, sheet=None):
    """Parse a workbook and describe what would be imported. Stores nothing."""
    try:
        import openpyxl
    except ImportError:
        return {"ok": False,
                "error": "openpyxl is not installed on this server, so workbook "
                         "uploads are unavailable. Enter the month by hand below, "
                         "or run: pip3.10 install --user openpyxl"}
    wb = openpyxl.load_workbook(path, data_only=True)
    idx = _alias_index(slug)
    events = parse_events(slug, wb)

    candidates = []
    for ws in wb.worksheets:
        row, cols = _find_header(ws)
        if not cols:
            continue
        label_col, hits = _find_label_col(ws, row, idx)
        candidates.append((hits, len(cols), ws.title, row, cols, label_col))
    if not candidates:
        return {"ok": False,
                "error": "No month headers found. The file needs a row of month "
                         "columns (e.g. Sep-26, Oct-26) and a column of line names."}

    if sheet:
        chosen = next((c for c in candidates if c[2] == sheet), None)
        if chosen is None:
            return {"ok": False, "error": f"Sheet '{sheet}' has no month headers."}
    else:
        chosen = max(candidates, key=lambda c: (c[0], c[1]))
    hits, _, title, header_row, month_cols, label_col = chosen
    ws = wb[title]

    base_months = {m["ym"]: m for m in store.baseline(slug)["months"]}
    in_model = {c: ym for c, ym in month_cols.items() if ym in base_months}

    rows, unmatched = [], []
    seen = set()
    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(r, label_col).value
        if raw is None or not str(raw).strip():
            continue
        key, conf = match_label(raw, idx)
        values = {}
        for c, ym in in_model.items():
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                values[ym] = float(v)
        if not key:
            if values:
                unmatched.append({"row": r, "raw": str(raw).strip()})
            continue
        if key in seen or not values:
            continue
        seen.add(key)
        rows.append({"row": r, "raw": str(raw).strip(), "key": key,
                     "label": next((ln["label"] for ln in analysis.line_meta(slug)
                                    if ln["key"] == key), str(raw).strip()),
                     "confidence": conf, "by_month": values})

    scale = _suggest_scale(rows, base_months, events)
    flip = _suggest_flip(rows, cost_keys(slug))

    # A running template carries a column for every month in the model. Only the
    # ones that actually hold a number are worth offering for import.
    with_data = {ym for row in rows for ym in row["by_month"]}
    if events:
        with_data |= set(events["months"])

    return {
        "ok": True,
        "sheet": title,
        "sheets": [c[2] for c in candidates],
        "header_row": header_row,
        "label_col": label_col,
        "months": sorted(with_data),
        "months_empty": sorted(set(in_model.values()) - with_data),
        "months_skipped": sorted({ym for ym in month_cols.values()
                                  if ym not in base_months}),
        "rows": rows,
        "unmatched": unmatched[:25],
        "suggest_scale": scale,
        "suggest_flip_costs": flip,
        "matched": len(rows),
        "events": events,
    }


def _suggest_scale(rows, base_months, events=None):
    """1 if the file is already in thousands, 0.001 if it is in whole dollars.

    Compares every dollar line against its own proforma value rather than a few
    named revenue lines -- on a template whose revenue lives on the Events sheet
    there may be no revenue line here at all, and the P&L sheet would then look
    like it was already in thousands.
    """
    ratios = []
    for row in rows:
        if row["key"] in NON_DOLLAR or row["key"].startswith("contracts__"):
            continue
        for ym, v in row["by_month"].items():
            plan = base_months.get(ym, {}).get("lines", {}).get(row["key"])
            if plan and abs(plan) > 1 and v:
                ratios.append(abs(v) / abs(plan))
    # The event list is a second, independent read on the units.
    if events:
        for slot in events["months"].values():
            for cats in slot["ancillary"].values():
                for payload in cats.values():
                    for k, v in payload.items():
                        if k.endswith("_rev") and v:
                            ratios.append(abs(v) / 10.0)
    if not ratios:
        return 1.0
    ratios.sort()
    median = ratios[len(ratios) // 2]
    return 0.001 if median > 100 else 1.0


def _suggest_flip(rows, costs):
    """True when cost lines arrive positive and need negating."""
    pos = neg = 0
    for row in rows:
        if row["key"] not in costs:
            continue
        for v in row["by_month"].values():
            if v > 0:
                pos += 1
            elif v < 0:
                neg += 1
    return pos > neg and pos > 0


def apply_preview(slug, prev, include_keys, months, scale=1.0, flip_costs=False,
                  filename=None):
    """Commit a confirmed preview into the actuals store, month by month."""
    include = set(include_keys)
    costs = cost_keys(slug)
    by_month, seg_by_month = {}, {}
    for row in prev["rows"]:
        if row["key"] not in include:
            continue
        # "contracts__barn" is a segment figure, not a venue line.
        seg_key = (row["key"].split("__", 1)[1]
                   if row["key"].startswith("contracts__") else None)
        for ym, v in row["by_month"].items():
            if ym not in months:
                continue
            if seg_key:
                seg_by_month.setdefault(ym, {}).setdefault(
                    seg_key, {})["contracts"] = v
                continue
            value = v
            if row["key"] not in NON_DOLLAR:
                value *= scale
                if flip_costs and row["key"] in costs and value > 0:
                    value = -value
            by_month.setdefault(ym, {})[row["key"]] = round(value, 6)

    written = []
    for ym in sorted(set(by_month) | set(seg_by_month)):
        store.record(slug, ym, by_month.get(ym, {}), "upload",
                     segments=seg_by_month.get(ym), note=filename)
        written.append(ym)
    return written


# ---------------------------------------------------------------- events sheet
# Attachment is far easier to report one event per row than as a month x category
# x metric grid, and an event list cross-checks the P&L revenue lines for free.
# Any sheet with a date column and at least two recognised category columns is
# treated as one.
EVENT_DATE_HEADERS = {"event date", "date", "event"}
EVENT_TYPE_HEADERS = {"contract type", "type", "contract", "cohort", "source",
                      "oc or new", "original or new"}
# Which venue or product line the event belongs to, where a site has more than
# one sharing a cost base.
EVENT_SEGMENT_HEADERS = {"venue", "segment", "site", "space", "location",
                         "venue space", "event type", "product"}
OC_WORDS = ("original", "oc", "assumed", "inherited", "pre close", "preclose",
            "backlog", "seller")

# Revenue Analysis category -> the names a reporting sheet is likely to use.
CATEGORY_ALIASES = {
    "room_rental": ["room rental", "venue fee", "venue rental", "rental", "site fee"],
    "food": ["food", "catering"],
    "beverage": ["beverage", "bar", "alcohol"],
    "lodging": ["lodging", "lodgings", "accommodation", "cabins", "overnight"],
    "dj": ["dj", "music", "entertainment"],
    "floral": ["floral", "flowers", "florals"],
    "bakery": ["bakery", "cake", "desserts"],
    "stationery": ["stationery", "stationary", "invitations", "paper"],
    "photography": ["photography", "photo", "photographer"],
    "other_ancillary": ["other ancillary", "other", "add ons", "enhancements",
                        "misc ancillary"],
    "service_charge": ["service charge", "service fee", "admin fee"],
    "elopement": ["elopement", "elopements", "elopement revenue"],
    "all_inclusive": ["all inclusive", "all-inclusive", "all inclusive package",
                      "package"],
}

EVENT_REVENUE_LINES = L.EVENT_REVENUE_LINES


def _category_index():
    idx = {}
    for key, names in CATEGORY_ALIASES.items():
        for name in names:
            idx.setdefault(norm(name), key)
    return idx


def _cohort(value):
    """'Original'/'OC'/'Inherited' mean the seller's backlog; anything else is new."""
    n = norm(value)
    return "oc" if any(w in n for w in OC_WORDS) else "new"


def find_events_sheet(wb):
    """(worksheet, header_row, {column: role}) for the first event-list sheet."""
    cat_idx = _category_index()
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, 25) + 1):
            roles, seen_cats = {}, 0
            for c in range(1, ws.max_column + 1):
                label = norm(ws.cell(r, c).value)
                if not label:
                    continue
                if label in EVENT_DATE_HEADERS and "date" not in roles.values():
                    roles[c] = "date"
                elif label in EVENT_TYPE_HEADERS and "type" not in roles.values():
                    roles[c] = "type"
                elif label in EVENT_SEGMENT_HEADERS and "segment" not in roles.values():
                    roles[c] = "segment"
                elif label in cat_idx:
                    roles[c] = cat_idx[label]
                    seen_cats += 1
            if "date" in roles.values() and seen_cats >= 2:
                return ws, r, roles
    return None, None, None


def _segment_index(slug):
    """Every spelling that should resolve to a segment key."""
    idx = {}
    for seg in store.baseline(slug).get("segments", []):
        for name in (seg["label"], seg["key"]):
            n = norm(name)
            idx.setdefault(n, seg["key"])
            idx.setdefault(n.rstrip("s"), seg["key"])
    return idx


def parse_events(slug, wb):
    """Roll an event list into per-month attachment counts, revenue and cohorts.

    Returns None when the workbook has no event sheet -- that is the normal case
    for a P&L-only file, not an error.
    """
    ws, header_row, roles = find_events_sheet(wb)
    if ws is None:
        return None

    base = store.baseline(slug)
    base_months = {m["ym"] for m in base["months"]}
    seg_idx = _segment_index(slug)
    segs = base.get("segments", [])
    default_seg = segs[0]["key"] if len(segs) == 1 else None

    date_col = next(c for c, role in roles.items() if role == "date")
    type_col = next((c for c, role in roles.items() if role == "type"), None)
    seg_col = next((c for c, role in roles.items() if role == "segment"), None)
    cat_cols = {c: role for c, role in roles.items()
                if role not in ("date", "type", "segment")}

    months, skipped, rows_read, unknown_segments = {}, [], 0, set()
    for r in range(header_row + 1, ws.max_row + 1):
        raw_date = ws.cell(r, date_col).value
        if not isinstance(raw_date, datetime):
            continue
        rows_read += 1
        ym = f"{raw_date.year:04d}-{raw_date.month:02d}"
        if ym not in base_months:
            skipped.append(ym)
            continue
        cohort = _cohort(ws.cell(r, type_col).value) if type_col else "new"

        seg_key = default_seg
        if seg_col is not None:
            raw_seg = ws.cell(r, seg_col).value
            n = norm(raw_seg)
            seg_key = seg_idx.get(n) or seg_idx.get(n.rstrip("s"))
            if seg_key is None and raw_seg is not None:
                unknown_segments.add(str(raw_seg).strip())
        if seg_key is None:
            continue

        slot = months.setdefault(ym, {"segments": {}, "ancillary": {}})
        seg_slot = slot["segments"].setdefault(
            seg_key, {"events_oc": 0, "events_new": 0, "events": 0})
        seg_slot[f"events_{cohort}"] += 1
        seg_slot["events"] += 1
        for c, cat in cat_cols.items():
            v = ws.cell(r, c).value
            if not isinstance(v, (int, float)) or isinstance(v, bool) or v <= 0:
                continue
            entry = slot["ancillary"].setdefault(seg_key, {}).setdefault(cat, {})
            entry[f"{cohort}_events"] = entry.get(f"{cohort}_events", 0) + 1
            entry[f"{cohort}_rev"] = entry.get(f"{cohort}_rev", 0.0) + float(v)

    return {
        "sheet": ws.title,
        "header_row": header_row,
        "categories": sorted(set(cat_cols.values())),
        "has_type_column": type_col is not None,
        "has_segment_column": seg_col is not None,
        "segments_needed": len(segs) > 1,
        "unknown_segments": sorted(unknown_segments),
        "rows_read": rows_read,
        "months": months,
        "months_skipped": sorted(set(skipped)),
    }


def apply_events(slug, parsed, months, scale=1.0, filename=None):
    """Commit parsed event detail. Counts are never scaled; money always is."""
    written = []
    present = set(parsed["categories"])
    for ym in sorted(parsed["months"]):
        if ym not in months:
            continue
        slot = parsed["months"][ym]

        # A category column that exists in the file but holds nothing means
        # nobody bought it -- a real zero. Left absent it would read as "not
        # reported" and the shortfall would disappear from the attachment tab.
        ancillary = {}
        for seg_key in slot["segments"]:
            base_cats = {cat: {"oc_events": 0, "new_events": 0,
                               "oc_rev": 0.0, "new_rev": 0.0} for cat in present}
            for cat, payload in slot["ancillary"].get(seg_key, {}).items():
                base_cats.setdefault(cat, {}).update(
                    {k: (v * scale if k.endswith("_rev") else v)
                     for k, v in payload.items()})
            ancillary[seg_key] = base_cats

        # Revenue lines the event list covers, summed across every segment --
        # the P&L is venue-wide even where the events are not.
        lines = {}
        for line, cats in EVENT_REVENUE_LINES.items():
            if not present.intersection(cats):
                continue
            total = 0.0
            for seg_cats in slot["ancillary"].values():
                for cat in cats:
                    payload = seg_cats.get(cat, {})
                    total += payload.get("oc_rev", 0.0) + payload.get("new_rev", 0.0)
            lines[line] = round(total * scale, 6)

        store.record(slug, ym, lines, "upload", ancillary=ancillary,
                     segments=slot["segments"], note=filename)
        written.append(ym)
    return written
